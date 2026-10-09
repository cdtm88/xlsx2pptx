"""Regenerate the spreadsheet fixtures used by tests/smoke.cjs.
Needs openpyxl (pip install openpyxl). Run from anywhere: python3 tests/fixtures/make.py"""
import datetime as dt
import os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
D = dt.datetime


def save(wb, name):
    wb.save(os.path.join(HERE, name))


# A three-task delivery plan in the current template layout.
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Plan"
ws.append(["Smoke Plan", None, None, None, None, None, None, None, "Start:", D(2026, 3, 4), "Now:", D(2026, 3, 18)])
ws.append([])
ws.append(["Task", "Category", "Status", "Start Date", "End Date", "Milestone", "At Risk", "Notes"])
for row in [["Alpha", "Build", "Done", D(2026, 3, 2), D(2026, 3, 13), False, False, None],
            ["Beta", "Build", "In Progress", D(2026, 3, 16), D(2026, 3, 27), False, True, None],
            ["Gamma", "Test", "Not Started", D(2026, 4, 1), D(2026, 4, 10), False, False, None]]:
    ws.append(row)
for r in ws.iter_rows(min_row=4):
    for c in r:
        if isinstance(c.value, dt.datetime):
            c.number_format = "dd/mm/yyyy"
save(wb, "delivery-plan.xlsx")

# A small Waqti timesheet export.
ART = "DN4 - dnata Travel Services"
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Export"
ws.append(["Staff ID", "Work Day", "Month", "Demand", "Feature", "Booked Hours", "Recharge Cost",
           "Application", "Taxonomy", "Staff ART", "Created By"])
for row in [["S1", "2026-03-02", "Mar-26", "D100 - Alpha", "F1 - Feature one", 6, 0, "App", "Dev", ART, "S1"],
            ["S2", "2026-03-03", "Mar-26", "D200 - Beta", "F2 - Feature two", 4, 0, "App", "Dev", ART, "S2"],
            ["S2", "2026-03-04", "Mar-26", "D200 - Beta", "F2 - Feature two", 2, 0, "", "Dev", ART, "GEN_LINES"]]:
    ws.append(row)
save(wb, "waqti-export.xlsx")
print("fixtures written to", HERE)
